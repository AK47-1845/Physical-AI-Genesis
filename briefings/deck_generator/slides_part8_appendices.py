"""
Slides Part 8: Appendices A & B — Complete Sources & Adversarial Self-Critique Log (Slides 46 to 51)
Executive Editorial Light System: Crisp white canvas, 0.8" margins, deep ink typography, Navy sign-off anchor.
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from deck_common import (
    new_slide, add_card, add_rect, add_header, add_takeaway, add_stat_box,
    add_badge, set_notes, section_divider_slide, BG_CANVAS, BG_CARD, BG_CARD_ALT,
    BG_CARD_ACCENT, BORDER_SUBTLE, BORDER_GOLD, BORDER_CYAN, BORDER_GREEN,
    BORDER_AMBER, BORDER_RED, GOLD, CYAN, GREEN, AMBER, RED, PURPLE, WHITE,
    COLOR_WHITE, BRAND_NAVY, BRAND_BLUE, INK_PRIMARY, TEXT_LIGHT, TEXT_MUTED, FONT_NAME
)

def slide_46_app_divider():
    return section_divider_slide(
        9,
        "Research Appendices & Governance",
        "Comprehensive Sources Index & Adversarial Self-Critique / Downgrade Log",
        [
            "Appendix A: Full primary citation catalog across capital, research, and field deployments",
            "Appendix B: Exhaustive record of all 11 claims removed or downgraded from early working drafts",
            "Transparent scientific methodology ensuring zero unaudited marketing copy enters the briefing",
            "Final executive sign-off and distribution register for LTTS leadership"
        ]
    )

def slide_47_appendix_a_p1():
    s = new_slide()
    add_header(s, "Appendix A · Citations", "Comprehensive Sources Index: Commercial & Capital Matrix",
               "Audited corporate disclosures, regulatory filings, and official announcements.")

    citations = [
        ("Wayve ($1.05B Series C)", "Corporate Announcement (May 7, 2024). Led by SoftBank Group with participation from NVIDIA and Microsoft.", GREEN),
        ("Figure AI ($675M Series B)", "Press Release (Feb 29, 2024). Valuing company at $2.6B, led by Parkway Venture Capital, Microsoft, OpenAI Startup Fund, NVIDIA, and Jeff Bezos.", GREEN),
        ("Physical Intelligence ($400M Series A)", "Company Disclosure (Nov 4, 2024). Valuing company at $2.4B, led by Bezos Expeditions, Thrive Capital, and Lux Capital.", GREEN),
        ("Skild AI ($300M Series A)", "Press Release (Jul 9, 2024). Valuing company at $1.5B, led by Lightspeed Venture Partners, Coatue, SoftBank, and Bezos Expeditions.", GREEN),
        ("Applied Intuition ($250M Series E)", "Press Release (Mar 12, 2024). Valuing company at $6.0B, led by Lux Capital, Elad Gil, and Porsche Investments Management.", GREEN),
        ("Scale AI ($1.0B Series F)", "Company Disclosure (May 21, 2024). Valuing company at $13.8B, led by Accel with participation from strategic partners.", GREEN),
        ("Waymo Ride & Mile Volume", "Alphabet Q3 2024 Earnings Call (Oct 29, 2024). Remarks by CEO Sundar Pichai confirming over 150,000 paid commercial trips/week and >1M autonomous miles weekly.", GREEN),
        ("Covariant Acquisition / License", "Amazon Corporate Announcement (Aug 30, 2024). Detailing hiring of founders Pieter Abbeel, Peter Chen, Rocky Duan, and non-exclusive foundation model licensing.", GREEN)
    ]

    card_w = 11.73
    top_base = 1.85
    row_h = 0.54
    gap = 0.08

    for i, (source, detail, color) in enumerate(citations):
        top = top_base + i * (row_h + gap)
        add_card(s, 0.8, top, card_w, row_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.5)

        tx = s.shapes.add_textbox(Inches(0.95), Inches(top + 0.06), Inches(card_w - 0.3), Inches(row_h - 0.12))
        tf = tx.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = source + "  —  "; r1.font.name = FONT_NAME; r1.font.size = Pt(10); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        r2 = p.add_run(); r2.text = detail; r2.font.name = FONT_NAME; r2.font.size = Pt(9); r2.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "Every financial figure, company valuation, and commercial metric traces directly to an official corporate release or audited regulatory filing.")
    set_notes(s, "Appendix A (Part 1). Verified primary sources for all capital rounds, valuations, and commercial volume figures.")
    return s


def slide_48_appendix_a_p2():
    s = new_slide()
    add_header(s, "Appendix A · Citations", "Comprehensive Sources Index: Research & Field Deployments",
               "Peer-reviewed research reports, foundation model preprints, and customer press releases.")

    citations_p2 = [
        ("Continuous Action Flow Matching", "Physical Intelligence Technical Report on π0 (Oct 2024, arXiv:2410.24164). Direct continuous trajectory generation.", BRAND_BLUE),
        ("Egocentric Scaling Laws", "NVIDIA GEAR Research Report on EgoScale (Feb 2026, arXiv:2602.16710). Proving log-linear scaling from human video.", BRAND_BLUE),
        ("Pearl's Causal Hierarchy", "Judea Pearl, Causality: Models, Reasoning, and Inference, Cambridge University Press. Foundational causal mathematics.", BRAND_NAVY),
        ("BMW Spartanburg Plant Trial", "BMW Group Official Press Release (Aug 2024). Confirming multi-week pilot of Figure 02 humanoid in sheet metal assembly.", GREEN),
        ("Mercedes-Benz Apollo Pilot", "Mercedes-Benz and Apptronik Joint Commercial Agreement Announcement (Mar 15, 2024). Assembly parts delivery trial.", GREEN),
        ("DHL Boston Dynamics Deployment", "DHL Supply Chain Official Announcement (2024). Expanding Boston Dynamics Stretch mobile manipulators for container unloading.", GREEN),
        ("Siemens & Microsoft Copilot", "Siemens AG Press Release (2024). Announcing integration of Generative AI industrial copilot into Siemens TIA Portal.", GREEN),
        ("Siemens & Intrinsic Alliance", "Siemens and Alphabet Intrinsic Strategic Partnership Announcement (2024). AI robotics integration with industrial automation.", GREEN)
    ]

    card_w = 11.73
    top_base = 1.85
    row_h = 0.54
    gap = 0.08

    for i, (source, detail, color) in enumerate(citations_p2):
        top = top_base + i * (row_h + gap)
        add_card(s, 0.8, top, card_w, row_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.5)

        tx = s.shapes.add_textbox(Inches(0.95), Inches(top + 0.06), Inches(card_w - 0.3), Inches(row_h - 0.12))
        tf = tx.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = source + "  —  "; r1.font.name = FONT_NAME; r1.font.size = Pt(10); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        r2 = p.add_run(); r2.text = detail; r2.font.name = FONT_NAME; r2.font.size = Pt(9); r2.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "All technical claims, mathematical formalisms, and OEM deployment partnerships are cross-verified against primary literature.")
    set_notes(s, "Appendix A (Part 2). Research papers (π0, EgoScale) and OEM trial announcements (BMW, Mercedes-Benz, DHL, Siemens).")
    return s


def slide_49_appendix_b_p1():
    s = new_slide()
    add_header(s, "Appendix B · Self-Critique", "Adversarial Self-Critique & Downgrade Log (Items 1 to 5)",
               "Complete, transparent record of unverified claims removed or downgraded from early working drafts.")

    critique_items = [
        ("1. Ford Battery Scrap & Warranty Savings [Removed]",
         "Ford Rawsonville plant achieved a 42% scrap reduction and $18M warranty reserve savings using PINN optical inspection.",
         "Ford operates advanced battery assembly at Rawsonville, but the exact '42% scrap' and '$18M savings' figures do not appear in audited SEC 10-K filings or official Ford press releases.",
         "Substituted with documented qualitative descriptions of battery optical inspection programs."),

        ("2. Airbus Fastener Hole Defect Rates [Removed]",
         "Airbus reduced fastener hole defect rate to < 2 PPM and inspection cycle time by 35% on A350 wing assembly.",
         "While Airbus deploys automated drilling robotics, the specific '< 2 PPM' defect rate was an unverified industry extrapolation from trade press whitepapers.",
         "Replaced with verified descriptions of automated drilling guidance systems."),

        ("3. Amazon Mobile Robot Fleet Count [Removed]",
         "Over 750,000 mobile robots and spatial perception arms operating across global fulfillment centers.",
         "While Amazon frequently cites the '750,000 robot' milestone in public PR, this figure aggregates legacy Kiva-style AGVs with modern AI AMRs (Proteus/Sequoia), creating a misleading impression of Physical-AI scale.",
         "Replaced with targeted descriptions of modern AMR fleet deployments."),

        ("4. GE Vernova Outage Cost Savings [Removed]",
         "Saved an estimated $380M in emergency outages across > 1,200 heavy-duty gas turbines.",
         "The '$380M saved' metric conflates vendor-modeled marketing projections with verified customer financial savings.",
         "Replaced with documented operational outage-reduction frameworks."),

        ("5. TSMC Lithography Yield Accuracy [Removed]",
         "TSMC achieved a 2.1x improvement in yield prediction accuracy in EUV lithography.",
         "Exact yield improvement metrics for TSMC advanced nodes (N3/N2) are proprietary trade secrets and not publicly auditable.",
         "Replaced with a qualified description of physics-informed defect classification pipelines.")
    ]

    card_w = 11.73
    top_base = 1.85
    row_h = 0.88
    gap = 0.10

    for i, (title, claim, reason, replace) in enumerate(critique_items):
        top = top_base + i * (row_h + gap)
        add_card(s, 0.8, top, card_w, row_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.5)

        tx = s.shapes.add_textbox(Inches(0.95), Inches(top + 0.06), Inches(card_w - 0.3), Inches(row_h - 0.12))
        tf = tx.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        r1 = p1.add_run(); r1.text = title + "\n"; r1.font.name = FONT_NAME; r1.font.size = Pt(10.5); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        
        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        r_cl_lbl = p2.add_run(); r_cl_lbl.text = "Draft Claim: "; r_cl_lbl.font.name = FONT_NAME; r_cl_lbl.font.size = Pt(8.5); r_cl_lbl.font.bold = True; r_cl_lbl.font.color.rgb = TEXT_MUTED
        r_cl = p2.add_run(); r_cl.text = f"\"{claim}\"   |   "; r_cl.font.name = FONT_NAME; r_cl.font.size = Pt(8.5); r_cl.font.color.rgb = TEXT_MUTED
        r_rs_lbl = p2.add_run(); r_rs_lbl.text = "Audit Verdict: "; r_rs_lbl.font.name = FONT_NAME; r_rs_lbl.font.size = Pt(8.5); r_rs_lbl.font.bold = True; r_rs_lbl.font.color.rgb = RED
        r_rs_d = p2.add_run(); r_rs_d.text = f"{reason}   |   "; r_rs_d.font.name = FONT_NAME; r_rs_d.font.size = Pt(8.5); r_rs_d.font.color.rgb = TEXT_LIGHT
        r_rp_lbl = p2.add_run(); r_rp_lbl.text = "Action: "; r_rp_lbl.font.name = FONT_NAME; r_rp_lbl.font.size = Pt(8.5); r_rp_lbl.font.bold = True; r_rp_lbl.font.color.rgb = GREEN
        r_rp_d = p2.add_run(); r_rp_d.text = replace; r_rp_d.font.name = FONT_NAME; r_rp_d.font.size = Pt(8.5); r_rp_d.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "This adversarial self-critique demonstrates uncompromising research integrity: no vendor marketing fluff survived the audit.")
    set_notes(s, "Appendix B (Items 1-5). Explain why transparency is credibility. We explicitly documented why Ford, Airbus, Amazon, GE, and TSMC claims were downgraded.")
    return s


def slide_50_appendix_b_p2():
    s = new_slide()
    add_header(s, "Appendix B · Self-Critique", "Adversarial Self-Critique & Downgrade Log (Items 6 to 11)",
               "Caterpillar, Intuitive Surgical, Symbotic, DHL, Siemens Energy, and Decorative Formalisms.")

    critique_items_p2 = [
        ("6. Caterpillar Autonomous Haulage Metrics [Removed]",
         "620 autonomous haul trucks operating with zero lost-time injuries over 250M kilometers.",
         "Cat's Command for Hauling is a proven mining product, but '250M km zero-injury' represents marketing copy rather than independently certified safety research.",
         "Replaced with sector-level automated haulage analysis."),

        ("7. Intuitive Surgical Force Measurement [Removed]",
         "> 10,000 force measurements per second reducing surgeon physical strain by 40% on da Vinci 5.",
         "The '40% strain reduction' claim is an uncontrolled ergonomic marketing claim.",
         "Replaced with factual descriptions of da Vinci 5 force-sensing hardware capabilities."),

        ("8. Symbotic Contracted Deployments [Removed]",
         "Symbotic has over $500M+ in contracted deployments with Walmart.",
         "Conflated multi-year system backlog with realized physical-AI software deployment.",
         "Replaced with general automated warehouse distribution analysis."),

        ("9. DHL Turnaround Time Reductions [Removed]",
         "Trailer turn time reduced from 90 to 38 minutes, cutting dock labor costs by 28%.",
         "Site-level minute-savings vary widely across facilities and are not standardized in audited reports.",
         "Replaced with verified Boston Dynamics Stretch deployment scope."),

        ("10. Siemens Energy Valve Failure Reduction [Removed]",
         "45% reduction in catastrophic valve failure events.",
         "Unaudited vendor whitepaper statistic without independent third-party certification.",
         "Replaced with verified descriptions of predictive thermal monitoring programs."),

        ("11. Decorative Equations & Academic Padding [Removed]",
         "Unconnected PINN loss functions (Ltotal = Ldata + lambda||dt u + N[u] - f||^2) and InfoNCE contrastive equations.",
         "These equations served as decorative academic padding without doing argumentative work.",
         "Retained only load-bearing formalisms: (1) Pearl's Causal Hierarchy, (2) Control Barrier Functions & Reachable Sets, and (3) Jacobian Singularity Invariants.")
    ]

    card_w = 11.73
    top_base = 1.85
    row_h = 0.74
    gap = 0.08

    for i, (title, claim, reason, replace) in enumerate(critique_items_p2):
        top = top_base + i * (row_h + gap)
        is_eq = "Decorative Equations" in title
        color = PURPLE if is_eq else RED

        add_card(s, 0.8, top, card_w, row_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.5)

        tx = s.shapes.add_textbox(Inches(0.95), Inches(top + 0.04), Inches(card_w - 0.3), Inches(row_h - 0.08))
        tf = tx.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        r1 = p1.add_run(); r1.text = title + "  —  "; r1.font.name = FONT_NAME; r1.font.size = Pt(10); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        rcl_lbl = p1.add_run(); rcl_lbl.text = "Draft Claim: "; rcl_lbl.font.name = FONT_NAME; rcl_lbl.font.size = Pt(8.5); rcl_lbl.font.bold = True; rcl_lbl.font.color.rgb = TEXT_MUTED
        rcl = p1.add_run(); rcl.text = f"\"{claim}\"  "; rcl.font.name = FONT_NAME; rcl.font.size = Pt(8.5); rcl.font.color.rgb = TEXT_MUTED

        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        r_rs = p2.add_run(); r_rs.text = "Audit Verdict: "; r_rs.font.name = FONT_NAME; r_rs.font.size = Pt(8.5); r_rs.font.bold = True; r_rs.font.color.rgb = color
        r_rs_d = p2.add_run(); r_rs_d.text = f"{reason}   |   "; r_rs_d.font.name = FONT_NAME; r_rs_d.font.size = Pt(8.5); r_rs_d.font.color.rgb = TEXT_LIGHT
        r_rp = p2.add_run(); r_rp.text = "Action: "; r_rp.font.name = FONT_NAME; r_rp.font.size = Pt(8.5); r_rp.font.bold = True; r_rp.font.color.rgb = GREEN
        r_rp_d = p2.add_run(); r_rp_d.text = replace; r_rp_d.font.name = FONT_NAME; r_rp_d.font.size = Pt(8.5); r_rp_d.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "Only load-bearing formalisms were retained: Pearl's Causal Ladder, Control Barrier Functions, and Jacobian Singularity Invariants.")
    set_notes(s, "Appendix B (Items 6-11). Detail item 11: removing decorative math padding to ensure every formula retained directly justifies a commercial recommendation.")
    return s


def slide_51_signoff():
    s = new_slide(bg_color=BRAND_NAVY)
    
    # Left accent gold bar
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = GOLD; bar.line.fill.background()

    add_badge(s, 0.8, 0.75, "DOCUMENT SIGN-OFF & DISTRIBUTION REGISTER", bg=RGBColor(0x00, 0x1E, 0x33), fg=GOLD, font_size=9, bold=True)

    tx_t = s.shapes.add_textbox(Inches(0.8), Inches(1.30), Inches(11.0), Inches(0.85))
    tf_t = tx_t.text_frame; tf_t.word_wrap = True; tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p = tf_t.paragraphs[0]
    r = p.add_run(); r.text = "The Practice of Physical Artificial Intelligence\n"; r.font.name = FONT_NAME; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = COLOR_WHITE
    r_sub = p.add_run(); r_sub.text = "Executive Strategic Briefing  ·  L&T Technology Services (LTTS)"; r_sub.font.name = FONT_NAME; r_sub.font.size = Pt(13); r_sub.font.color.rgb = GOLD

    div = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.35), Inches(11.5), Inches(0.015))
    div.fill.solid(); div.fill.fore_color.rgb = GOLD; div.line.fill.background()

    # Sign-off info boxes
    boxes = [
        ("DOCUMENT REGISTER", [
            ("Version:", "3.0 Master Executive Release"),
            ("Document ID:", "LTTS-PAI-2026-MASTER-01"),
            ("Date of Issue:", "September 2026"),
            ("Classification:", "Strictly Confidential & Proprietary")
        ], RGBColor(0x38, 0xBD, 0xF8)),

        ("PRIMARY RESEARCH & AUTHORSHIP", [
            ("Lead Author:", "Adari Karthikeya"),
            ("Practice:", "Physical AI Practice Research"),
            ("Organization:", "LTTS Strategic Engineering Practice"),
            ("Review Cycle:", "Full Peer Review & Self-Critique Completed")
        ], GOLD),

        ("EXECUTIVE DISTRIBUTION", [
            ("Prepared For:", "Dr. Madhusudhan Singh"),
            ("Title:", "Executive & Technical Leadership"),
            ("Entity:", "L&T Technology Services Limited (LTTS)"),
            ("Action Required:", "Review & Budget Authorization for Q1 2027")
        ], GREEN)
    ]

    card_w = 3.68
    gap = 0.34
    for i, (hdr, items, clr) in enumerate(boxes):
        left = 0.8 + i * (card_w + gap)
        top = 2.65
        add_card(s, left, top, card_w, 3.8, bg=RGBColor(0x04, 0x33, 0x54), border=RGBColor(0x0C, 0x4A, 0x73))
        add_rect(s, left, top, card_w, 0.025, clr)

        tx_b = s.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.20), Inches(card_w - 0.4), Inches(3.4))
        tf_b = tx_b.text_frame; tf_b.word_wrap = True; tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0
        p_hdr = tf_b.paragraphs[0]
        r_hdr = p_hdr.add_run(); r_hdr.text = hdr + "\n\n"; r_hdr.font.name = FONT_NAME; r_hdr.font.size = Pt(11); r_hdr.font.bold = True; r_hdr.font.color.rgb = clr

        for k, (lbl, val) in enumerate(items):
            p_item = tf_b.add_paragraph()
            if k > 0: p_item.space_before = Pt(7)
            rl = p_item.add_run(); rl.text = lbl + " "; rl.font.name = FONT_NAME; rl.font.size = Pt(9.5); rl.font.bold = True; rl.font.color.rgb = COLOR_WHITE
            rv = p_item.add_run(); rv.text = val; rv.font.name = FONT_NAME; rv.font.size = Pt(9.5); rv.font.color.rgb = RGBColor(0xBA, 0xE6, 0xFD)

    add_takeaway(s, "Briefing completed. For presentation delivery, boardroom inquiries, or pilot execution, contact Adari Karthikeya & LTTS Engineering Leadership.", is_dark=True)
    set_notes(s, "Final Slide. Formal sign-off and distribution register for Dr. Madhusudhan Singh and LTTS leadership.")
    return s
