"""
Slides Part 4: Module 3 — Commercialization & LTTS Playbook (Slides 18 to 22)
Executive Editorial Light System: Crisp white canvas, 0.8" margins, deep ink typography.
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from deck_common import (
    new_slide, add_card, add_rect, add_header, add_takeaway, add_stat_box,
    add_badge, set_notes, section_divider_slide, BG_CANVAS, BG_CARD, BG_CARD_ALT,
    BG_CARD_ACCENT, BORDER_SUBTLE, BORDER_GOLD, BORDER_CYAN, BORDER_GREEN,
    BORDER_AMBER, BORDER_RED, BORDER_PURPLE, GOLD, CYAN, GREEN, AMBER, RED, PURPLE, WHITE,
    COLOR_WHITE, BRAND_NAVY, BRAND_BLUE, INK_PRIMARY, TEXT_LIGHT, TEXT_MUTED, FONT_NAME
)

def slide_18_m3_divider():
    return section_divider_slide(
        3,
        "Commercialization Dynamics",
        "How Physical AI Generates Revenue & The LTTS Engineering Services Playbook",
        [
            "Commercial models, contract structures, ACV ranges, and sales cycles (Table 3)",
            "Primary buyer personas across robotics R&D, plant operations, and safety",
            "The engineering services opportunity: addressing client deficits without hardware capital drag",
            "LTTS's four commercial offerings: SD-FaaS, PLC Integration, V&V, and Embodied MLOps"
        ]
    )

def slide_19_commercial_models_table():
    s = new_slide()
    add_header(s, "Module 3 · Commercialization", "Commercial Models, Deal Economics & Buyer Personas (Table 3)",
               "Comprehensive benchmark of revenue archetypes across foundation models, OEMs, software, and systems integration.")

    # Executive Table Container Layout
    left_m = 0.8
    top_base = 1.80
    total_w = 11.73

    # Column widths summing to 11.73"
    col_w = [2.70, 2.45, 2.90, 1.38, 2.30]
    col_x = [left_m]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)

    # Header Row (Authoritative Deep Navy)
    hdr_h = 0.38
    add_rect(s, left_m, top_base, total_w, hdr_h, BRAND_NAVY)

    headers = [
        "BUSINESS MODEL ARCHETYPE",
        "REPRESENTATIVE ENTITIES",
        "CONTRACT & DEAL ECONOMICS",
        "SALES CYCLE",
        "PRIMARY BUYER PERSONA"
    ]

    for j, (h_text, x_pos, w) in enumerate(zip(headers, col_x, col_w)):
        tx = s.shapes.add_textbox(Inches(x_pos + 0.10), Inches(top_base + 0.05), Inches(w - 0.20), Inches(hdr_h - 0.10))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = h_text
        r.font.name = FONT_NAME
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_WHITE

    # Table Data Rows
    models = [
        ("Model Licensing / FMaaS",
         "Physical Intelligence, Skild AI, Covariant",
         "Annual base license + API overage\n($100k - $500k+ ARR) [Estimated]",
         "3 - 6 mos",
         "VP Software /\nHead of Robotics R&D",
         CYAN, False),

        ("Integrated HW+SW OEM",
         "Figure AI, Agility Robotics, Boston Dynamics",
         "Robotics-as-a-Service ($5k-$15k/robot/mo)\nor CapEx + maintenance SLA [Reported]",
         "12 - 18 mos",
         "VP Logistics /\nPlant General Manager",
         AMBER, False),

        ("Simulation & Validation Tooling",
         "NVIDIA Isaac Sim, Applied Intuition",
         "Enterprise SaaS subscription\n($250k - $2.0M+ ACV) [Reported]",
         "6 - 9 mos",
         "VP Engineering /\nHead of Autonomous Systems",
         GREEN, False),

        ("Compliance & Safety Verification",
         "TÜV SÜD, UL Solutions, Credo AI",
         "Safety audit fee + recurring testing suite\n($150k - $750k per engagement) [Estimated]",
         "6 - 12 mos",
         "Chief Safety Officer /\nVP Quality & Compliance",
         PURPLE, False),

        ("Systems Integration & Deployment",
         "LTTS, KPIT Tech, Accenture Industry X",
         "T&M + Milestone-gated deployment\n($1.0M - $8.0M+ per site) [Estimated]",
         "4 - 8 mos",
         "Chief Information Officer /\nVP Manufacturing Operations",
         BRAND_BLUE, True)  # Highlighted LTTS Core Play
    ]

    row_h = 0.88
    for i, (bmodel, entities, contract, cycle, buyer, accent_c, is_ltts) in enumerate(models):
        row_top = top_base + hdr_h + i * row_h
        row_bg = RGBColor(0xEE, 0xF4, 0xFA) if is_ltts else (BG_CARD if i % 2 == 0 else COLOR_WHITE)
        row_border = BRAND_BLUE if is_ltts else BORDER_SUBTLE

        # Row Background Card
        add_card(s, left_m, row_top, total_w, row_h, bg=row_bg, border=row_border, border_width=1.0 if is_ltts else 0.5)

        if is_ltts:
            add_rect(s, left_m, row_top, 0.04, row_h, BRAND_BLUE)

        # Col 0: Business Model Archetype
        tx0 = s.shapes.add_textbox(Inches(col_x[0] + 0.12), Inches(row_top + 0.12), Inches(col_w[0] - 0.20), Inches(row_h - 0.24))
        tf0 = tx0.text_frame; tf0.word_wrap = True; tf0.margin_left = tf0.margin_right = tf0.margin_top = tf0.margin_bottom = 0
        p0 = tf0.paragraphs[0]
        r0 = p0.add_run(); r0.text = bmodel; r0.font.name = FONT_NAME; r0.font.size = Pt(10.5); r0.font.bold = True; r0.font.color.rgb = INK_PRIMARY
        if is_ltts:
            p0_b = tf0.add_paragraph()
            p0_b.space_before = Pt(2)
            r0_b = p0_b.add_run(); r0_b.text = "★ LTTS CORE OPPORTUNITY"; r0_b.font.name = FONT_NAME; r0_b.font.size = Pt(8); r0_b.font.bold = True; r0_b.font.color.rgb = BRAND_BLUE

        # Col 1: Representative Entities
        tx1 = s.shapes.add_textbox(Inches(col_x[1] + 0.10), Inches(row_top + 0.14), Inches(col_w[1] - 0.20), Inches(row_h - 0.28))
        tf1 = tx1.text_frame; tf1.word_wrap = True; tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0
        p1 = tf1.paragraphs[0]
        r1 = p1.add_run(); r1.text = entities; r1.font.name = FONT_NAME; r1.font.size = Pt(9); r1.font.color.rgb = TEXT_LIGHT

        # Col 2: Contract Structure & Size
        tx2 = s.shapes.add_textbox(Inches(col_x[2] + 0.10), Inches(row_top + 0.12), Inches(col_w[2] - 0.20), Inches(row_h - 0.24))
        tf2 = tx2.text_frame; tf2.word_wrap = True; tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        r2 = p2.add_run(); r2.text = contract; r2.font.name = FONT_NAME; r2.font.size = Pt(9); r2.font.color.rgb = TEXT_LIGHT

        # Col 3: Sales Cycle
        tx3 = s.shapes.add_textbox(Inches(col_x[3] + 0.08), Inches(row_top + 0.20), Inches(col_w[3] - 0.16), Inches(row_h - 0.32))
        tf3 = tx3.text_frame; tf3.word_wrap = True; tf3.margin_left = tf3.margin_right = tf3.margin_top = tf3.margin_bottom = 0
        p3 = tf3.paragraphs[0]
        r3 = p3.add_run(); r3.text = cycle; r3.font.name = FONT_NAME; r3.font.size = Pt(9.5); r3.font.bold = True
        r3.font.color.rgb = GREEN if "4 - 8" in cycle or "3 - 6" in cycle else TEXT_LIGHT

        # Col 4: Primary Buyer Persona
        tx4 = s.shapes.add_textbox(Inches(col_x[4] + 0.10), Inches(row_top + 0.12), Inches(col_w[4] - 0.20), Inches(row_h - 0.24))
        tf4 = tx4.text_frame; tf4.word_wrap = True; tf4.margin_left = tf4.margin_right = tf4.margin_top = tf4.margin_bottom = 0
        p4 = tf4.paragraphs[0]
        r4 = p4.add_run(); r4.text = buyer; r4.font.name = FONT_NAME; r4.font.size = Pt(9); r4.font.bold = is_ltts; r4.font.color.rgb = INK_PRIMARY

    add_takeaway(s, "Systems integration captures the largest single-site budget ($1.0M-$8.0M+) with relatively fast 4-8 month enterprise sales cycles.")
    set_notes(s, "Review Table 3. Point out that while hardware OEMs face 12-18 month sales cycles, systems integrators like LTTS operate on 4-8 month cycles with high deal sizes.")
    return s


def slide_20_engineering_services_opportunity():
    s = new_slide()
    add_header(s, "Module 3 · Commercialization", "The Engineering Services Opportunity: The LTTS Advantage",
               "Solving the structural capability deficit of industrial clients without taking on hardware capital drag.")

    # 3 Stat Cards on Top
    add_stat_box(s, 0.8, 1.80, 3.75, 1.25, "$1.0M-$8.0M+", "Per-Site Deployment", "Systems integration & brownfield rollout budget [Estimated]", BRAND_NAVY)
    add_stat_box(s, 4.8, 1.80, 3.75, 1.25, "Zero Hardware", "Capital Drag Avoided", "No robotics manufacturing or warranty liability [Verified]", GREEN)
    add_stat_box(s, 8.8, 1.80, 3.73, 1.25, "4-8 Months", "Rapid Sales Velocity", "Direct sponsorship from enterprise CIOs & VP Mfg [Verified]", BRAND_BLUE)

    # Left: The Industrial Deficit
    add_card(s, 0.8, 3.25, 5.7, 3.35, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.75)
    add_rect(s, 0.8, 3.25, 5.7, 0.025, AMBER)
    tx_l = s.shapes.add_textbox(Inches(1.05), Inches(3.45), Inches(5.2), Inches(3.0))
    tf_l = tx_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0

    p1 = tf_l.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "THE CLIENT CAPABILITY DEFICIT\n"
    r1.font.name = FONT_NAME; r1.font.size = Pt(11); r1.font.bold = True; r1.font.color.rgb = AMBER

    points_def = [
        ("Brownfield Plant Heterogeneity:", "Automotive and industrial plants run decades-old Siemens, Rockwell, and Mitsubishi PLCs. They cannot simply plug in Python APIs from AI startups."),
        ("Compliance & Liability Gaps:", "Plant managers cannot deploy uncertified neural policies without voiding insurance or risking catastrophic downtime."),
        ("CAD-to-Sim Bottlenecks:", "Industrial clients own immense CAD and PLM assets (Siemens Teamcenter, NX, CATIA) but lack the internal capability to convert them into simulation digital twins.")
    ]
    for h, d in points_def:
        p = tf_l.add_paragraph()
        p.space_before = Pt(5)
        rh = p.add_run()
        rh.text = f"•  {h} "
        rh.font.name = FONT_NAME; rh.font.size = Pt(9.5); rh.font.bold = True; rh.font.color.rgb = INK_PRIMARY
        rd = p.add_run()
        rd.text = d
        rd.font.name = FONT_NAME; rd.font.size = Pt(9); rd.font.color.rgb = TEXT_LIGHT

    # Right: LTTS's Positioning
    add_card(s, 6.83, 3.25, 5.7, 3.35, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.75)
    add_rect(s, 6.83, 3.25, 5.7, 0.025, GREEN)
    tx_r = s.shapes.add_textbox(Inches(7.08), Inches(3.45), Inches(5.2), Inches(3.0))
    tf_r = tx_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = 0

    pr1 = tf_r.paragraphs[0]
    rr1 = pr1.add_run()
    rr1.text = "LTTS'S STRATEGIC PLAYBOOK ADVANTAGE\n"
    rr1.font.name = FONT_NAME; rr1.font.size = Pt(11); rr1.font.bold = True; rr1.font.color.rgb = GREEN

    points_adv = [
        ("Pure-Play Engineering Heritage:", "LTTS possesses decades of deep industrial automation, mechatronics, and safety engineering DNA that pure-software AI labs completely lack."),
        ("Trusted Plant Access:", "Already embedded across hundreds of Tier-1 industrial manufacturing plants worldwide with established master service agreements (MSAs)."),
        ("The Neutral Bridge:", "LTTS does not compete with robot OEMs or foundation model builders. LTTS partners with all of them, capturing services revenue from each deployment.")
    ]
    for h, d in points_adv:
        p = tf_r.add_paragraph()
        p.space_before = Pt(5)
        rh = p.add_run()
        rh.text = f"✔  {h} "
        rh.font.name = FONT_NAME; rh.font.size = Pt(9.5); rh.font.bold = True; rh.font.color.rgb = INK_PRIMARY
        rd = p.add_run()
        rd.text = d
        rd.font.name = FONT_NAME; rd.font.size = Pt(9); rd.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "LTTS bridges the massive gap between agile Silicon Valley AI models and rigid, highly regulated brownfield plant operations.")
    set_notes(s, "The Engineering Services Opportunity. Explain why pure software startups fail in factories: they don't understand fieldbuses, safety PLCs, or ISO standards.")
    return s


def slide_21_offerings_1_and_2():
    s = new_slide()
    add_header(s, "Module 3 · Commercialization", "LTTS Commercial Offerings (1 & 2): SD-FaaS & PLC Integration",
               "High-margin technical services transforming client digital assets and plant automation.")

    offerings = [
        ("OFFERING 1", "Synthetic Data Factory as a Service (SD-FaaS)", BRAND_NAVY, BORDER_SUBTLE,
         "$1.0M - $3.5M Annual Recurring Engagements [Estimated]",
         "Target Client: Automotive, Aerospace, and Industrial Machinery OEMs",
         [
             ("Asset Ingestion:", "Ingesting complex client CAD, BIM, and PLM assemblies directly from Siemens NX, Teamcenter, CATIA, and PTC Creo."),
             ("Digital Twin Construction:", "Automating the conversion of mechanical models into kinematic, contact-calibrated digital twin environments in NVIDIA Isaac Sim and Omniverse."),
             ("Extreme Domain Randomization:", "Generating millions of synthetic photorealistic frames and physics variations (lighting, textures, friction, tolerances) to train robust robot policies."),
             ("Monetization Model:", "Annual platform subscription + fee per thousand validated scenario hours.")
         ]),

        ("OFFERING 2", "Brownfield PLC-to-AI Integration Middleware", BRAND_BLUE, BORDER_CYAN,
         "$1.5M - $5.0M Per Plant Rollout [Estimated]",
         "Target Client: Discrete Manufacturing, Tier-1 Auto Suppliers, Warehouses",
         [
             ("Deterministic Bridge Middleware:", "Engineering hard-real-time software bridges connecting high-level neural policy outputs to industrial field controllers."),
             ("Fieldbus Protocol Interfacing:", "Zero-jitter integration with Siemens S7-1500 and Rockwell ControlLogix PLCs via Profinet, EtherCAT, and OPC UA at sub-millisecond latency."),
             ("Hardware Safety Interlocks:", "Embedding deterministic safety envelopes that cut motor power via Cat 4 / PL e safety relays if neural policies violate velocity or keep-out zones."),
             ("Monetization Model:", "Base integration fee per cell + site license + SLA maintenance.")
         ])
    ]

    card_w = 5.7
    for i, (badge, title, color, border, rev, client, items) in enumerate(offerings):
        left = 0.8 + i * (card_w + 0.33)
        top = 1.85
        card_h = 4.75

        add_card(s, left, top, card_w, card_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.75)
        add_rect(s, left, top, card_w, 0.025, color)

        add_badge(s, left + 0.25, top + 0.22, badge, bg=border, fg=color, font_size=8.5, bold=True)

        tx = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.58), Inches(card_w - 0.5), Inches(0.85))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = title + "\n"
        r1.font.name = FONT_NAME; r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        r2 = p.add_run()
        r2.text = rev
        r2.font.name = FONT_NAME; r2.font.size = Pt(9.5); r2.font.bold = True; r2.font.color.rgb = color

        tx_c = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 1.48), Inches(card_w - 0.5), Inches(0.35))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        pc = tf_c.paragraphs[0]
        rc = pc.add_run()
        rc.text = client
        rc.font.name = FONT_NAME; rc.font.size = Pt(8.8); rc.font.italic = True; rc.font.color.rgb = TEXT_MUTED

        tx_b = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 1.90), Inches(card_w - 0.5), Inches(2.7))
        tf_b = tx_b.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0

        for j, (hdr, desc) in enumerate(items):
            pj = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            if j > 0: pj.space_before = Pt(5)
            rj1 = pj.add_run()
            rj1.text = f"•  {hdr} "
            rj1.font.name = FONT_NAME; rj1.font.size = Pt(9.5); rj1.font.bold = True; rj1.font.color.rgb = INK_PRIMARY
            rj2 = pj.add_run()
            rj2.text = desc
            rj2.font.name = FONT_NAME; rj2.font.size = Pt(9); rj2.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "SD-FaaS monetizes enterprise CAD assets into synthetic data; Brownfield Integration secures plant floor PLC fieldbus control.")
    set_notes(s, "Detail Offerings 1 and 2. Emphasize that clients already have massive CAD files (Siemens Teamcenter); LTTS converts them into synthetic data pipelines.")
    return s


def slide_22_offerings_3_and_4():
    s = new_slide()
    add_header(s, "Module 3 · Commercialization", "LTTS Commercial Offerings (3 & 4): V&V Practice & Embodied MLOps",
               "Regulatory safety assurance and continuous lifecycle maintenance across plant operations.")

    offerings = [
        ("OFFERING 3", "Physical AI Verification & Validation (V&V)", GREEN, BORDER_GREEN,
         "$500k - $1.5M Per Engagement [Estimated]",
         "Target Client: Regulated OEMs, Factory Operators, Model Developers",
         [
             ("Independent Neural Stress-Testing:", "Independent rigorous auditing and stress-testing of third-party neural controllers and VLAs under edge-case conditions."),
             ("Regulatory Compliance Audits:", "Certifying neural controllers against ISO 13849 (Machinery Safety) and IEC 61508 (Functional Safety) international standards."),
             ("Automated Safety Dossier Compiler:", "Software toolchain evaluating Control Barrier Functions (CBFs) and forward reachable sets via interval bound propagation (alpha,beta-CROWN)."),
             ("Monetization Model:", "Fixed audit engagement fee + compliance dossier certification report.")
         ]),

        ("OFFERING 4", "Managed Embodied MLOps & Fleet Telemetry", PURPLE, BORDER_PURPLE,
         "$500k - $2.0M Annual Recurring Revenue (ARR) [Estimated]",
         "Target Client: Enterprise Robot Fleet Operators, Logistics Hubs",
         [
             ("Continuous Telemetry Logging:", "Real-time edge ingestion of motor torques, currents, joint temperatures, and visual odometry across deployed fleets."),
             ("Automated Drift Detection:", "Automated monitoring detecting model drift, mechanical wear, friction changes, or out-of-distribution environmental anomalies."),
             ("Automated Retraining Loop:", "Flagged edge anomalies automatically trigger synthetic re-simulation in Isaac Sim, continuous fine-tuning, and OTA policy deployment."),
             ("Monetization Model:", "Per-robot/per-cell annual managed services subscription with 99.9% uptime SLA.")
         ])
    ]

    card_w = 5.7
    for i, (badge, title, color, border, rev, client, items) in enumerate(offerings):
        left = 0.8 + i * (card_w + 0.33)
        top = 1.85
        card_h = 4.75

        add_card(s, left, top, card_w, card_h, bg=BG_CARD, border=BORDER_SUBTLE, border_width=0.75)
        add_rect(s, left, top, card_w, 0.025, color)

        add_badge(s, left + 0.25, top + 0.22, badge, bg=border, fg=color, font_size=8.5, bold=True)

        tx = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.58), Inches(card_w - 0.5), Inches(0.85))
        tf = tx.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = title + "\n"
        r1.font.name = FONT_NAME; r1.font.size = Pt(12.5); r1.font.bold = True; r1.font.color.rgb = INK_PRIMARY
        r2 = p.add_run()
        r2.text = rev
        r2.font.name = FONT_NAME; r2.font.size = Pt(9.5); r2.font.bold = True; r2.font.color.rgb = color

        tx_c = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 1.48), Inches(card_w - 0.5), Inches(0.35))
        tf_c = tx_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
        pc = tf_c.paragraphs[0]
        rc = pc.add_run()
        rc.text = client
        rc.font.name = FONT_NAME; rc.font.size = Pt(8.8); rc.font.italic = True; rc.font.color.rgb = TEXT_MUTED

        tx_b = s.shapes.add_textbox(Inches(left + 0.25), Inches(top + 1.90), Inches(card_w - 0.5), Inches(2.7))
        tf_b = tx_b.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = tf_b.margin_bottom = 0

        for j, (hdr, desc) in enumerate(items):
            pj = tf_b.paragraphs[0] if j == 0 else tf_b.add_paragraph()
            if j > 0: pj.space_before = Pt(5)
            rj1 = pj.add_run()
            rj1.text = f"•  {hdr} "
            rj1.font.name = FONT_NAME; rj1.font.size = Pt(9.5); rj1.font.bold = True; rj1.font.color.rgb = INK_PRIMARY
            rj2 = pj.add_run()
            rj2.text = desc
            rj2.font.name = FONT_NAME; rj2.font.size = Pt(9); rj2.font.color.rgb = TEXT_LIGHT

    add_takeaway(s, "V&V establishes LTTS as the trusted certification gatekeeper, while Managed MLOps generates long-term recurring SaaS/services ARR.")
    set_notes(s, "Detail Offerings 3 and 4. V&V is an immediate high-margin revenue generator (Q1 2027); MLOps establishes permanent client lock-in.")
    return s
